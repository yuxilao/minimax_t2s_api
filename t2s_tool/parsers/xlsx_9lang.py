from __future__ import annotations

import os
import re

import openpyxl

from ..errors import FileError, ParseError
from ..models import LanguagePack, ParseResult, VoiceEntry
from . import register_parser
from .base import ParserBase

SKIP_COLUMNS = {"AP音效名称", "场景详细描述"}
ID_RE = re.compile(r"^Q\d+$")


@register_parser
class Xlsx9LangParser(ParserBase):
    name = "xlsx_9lang"
    display_name = "9国语种语音列表(xlsx)"
    extensions = (".xlsx",)

    def detect(self, path: str) -> int:
        if not str(path).lower().endswith(".xlsx") or not os.path.exists(path):
            return 0
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            try:
                ws = wb[wb.sheetnames[0]]
                headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            finally:
                wb.close()
        except Exception:
            return 0
        if not headers:
            return 0
        langs = [h for h in headers[1:] if h and str(h).strip() not in SKIP_COLUMNS]
        return 80 if len(langs) >= 2 else 0

    def parse(self, path: str) -> ParseResult:
        if not os.path.exists(path):
            raise FileError(f"语音目录文件不存在: {path}")
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
        except Exception as e:
            raise ParseError(f"无法打开 Excel 文件: {e}")
        warnings = []
        try:
            ws = wb[wb.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                raise ParseError("Excel 首行(表头)为空")
            lang_cols = [
                (i, str(h).strip())
                for i, h in enumerate(headers)
                if i > 0 and h and str(h).strip() and str(h).strip() not in SKIP_COLUMNS
            ]
            if not lang_cols:
                raise ParseError("未找到任何语言列（表头除 AP音效名称/场景详细描述 外为空）")
            buckets = {h: [] for _, h in lang_cols}
            seen = set()
            for rno, row in enumerate(rows, 2):
                vid = str(row[0]).strip() if row and row[0] is not None else ""
                if not vid:
                    continue  # 空行跳过
                if not ID_RE.match(vid):
                    if any(c is not None and str(c).strip() for c in row[1:]):
                        warnings.append(f"第{rno}行语音ID非法，已跳过: {vid}")
                    continue
                if vid in seen:
                    warnings.append(f"重复语音ID已去重: {vid}")
                    continue
                seen.add(vid)
                for i, h in lang_cols:
                    raw = row[i] if i < len(row) else None
                    text = str(raw).replace("\n", " ").replace("\r", "").strip() if raw is not None else ""
                    if not text:
                        warnings.append(f"{vid} 的「{h}」为空文本")
                    buckets[h].append(VoiceEntry(vid, text))
        finally:
            wb.close()
        packs = [
            LanguagePack(h, sorted(es, key=lambda e: int(e.voice_id[1:])))
            for h, es in buckets.items()
        ]
        return ParseResult(path, self.name, packs, warnings)
